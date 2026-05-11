from pathlib import Path

import xlrd

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook as OpenPyxlWorkbook

"""
    Faz o carregamento do arquivo excel independente se é um XLS ou XLSX
    e retorna um objeto openpyxl.workbook.
    Motivo: A biblioteca "openpyxl" não lê arquivos antigos tipo xls, então quando é xls o arquivo é aberto pela biblioteca "xlrd"
    e depois lida e copiada para um objeto openpyxl.sheet e retornado o objeto (virtual, pois não é criado um arquivo físico).
    OBSERVAÇÃO: está copia é limitada e contém apenas texto, ignorando imagens, formatações, etc. mas serve para extrair os dados do arquivo xls. 
"""
class ExcelLoader:

    def load(self, path: str | Path) -> OpenPyxlWorkbook:

        # Converte para objeto Path
        path = Path(path)

        # Valida existência do arquivo
        if not path.exists():

            raise FileNotFoundError(
                f"Arquivo não encontrado: {path}"
            )

        # Obtém extensão do arquivo
        extension = path.suffix.lower()

        # Carrega arquivo .xlsx diretamente
        if extension == ".xlsx":

            return self._load_xlsx(path)

        # Carrega arquivo .xls convertendo para openpyxl
        if extension == ".xls":

            return self._load_xls(path)

        # Extensão inválida
        raise ValueError(
            f"Formato de arquivo não suportado: {extension}"
        )

    def _load_xlsx(
        self,
        path: Path
    ) -> OpenPyxlWorkbook:

        # Carrega workbook padrão openpyxl
        return load_workbook(
            filename=path,
            data_only=False
        )

    def _load_xls(self, path: Path) -> OpenPyxlWorkbook:

        # Abre workbook .xls usando xlrd
        xls_workbook = xlrd.open_workbook(path)

        # Cria workbook openpyxl vazio
        workbook = Workbook()

        # Remove worksheet padrão criada automaticamente
        workbook.remove(workbook.active)

        # Percorre sheets do arquivo .xls
        for xls_sheet in xls_workbook.sheets():

            # Cria worksheet equivalente no openpyxl
            worksheet = workbook.create_sheet(
                title=xls_sheet.name
            )

            # Copia conteúdo das células
            self._copy_sheet_data(
                xls_sheet=xls_sheet,
                worksheet=worksheet
            )

            # Copia merges existentes
            self._copy_merged_cells(
                xls_sheet=xls_sheet,
                worksheet=worksheet
            )

        # Define primeira worksheet como ativa
        if workbook.worksheets:
            workbook.active = 0

        # Retorna workbook convertido
        return workbook

    def _copy_sheet_data(self, xls_sheet, worksheet) -> None:

        # Percorre linhas da sheet
        for row in range(xls_sheet.nrows):

            # Percorre colunas da sheet
            for col in range(xls_sheet.ncols):

                # Obtém valor da célula
                value = xls_sheet.cell_value(
                    row,
                    col
                )

                # Cria célula equivalente no openpyxl
                worksheet.cell(
                    row=row + 1,
                    column=col + 1,
                    value=value
                )

    def _copy_merged_cells(self, xls_sheet, worksheet) -> None:

        # Percorre ranges merged do xlrd
        for merged_range in xls_sheet.merged_cells:

            (
                row_start,
                row_end,
                col_start,
                col_end
            ) = merged_range

            # Ajusta índices para padrão openpyxl
            start_row = row_start + 1
            end_row = row_end

            start_col = col_start + 1
            end_col = col_end

            # Recria merge no worksheet openpyxl
            worksheet.merge_cells(
                start_row=start_row,
                end_row=end_row,
                start_column=start_col,
                end_column=end_col
            )