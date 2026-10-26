from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


class ExcelExport:

    @staticmethod
    def export(rows, headers, file_prefix):
        export_dir = Path("exportacoes")
        export_dir.mkdir(exist_ok=True)

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        file_path = export_dir / f"{file_prefix}_{timestamp}.xlsx"

        wb = Workbook()
        ws = wb.active

        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in rows:
            ws.append(row)

        if file_prefix == "pesquisa_por_dia_e_entrada":
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)

            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=ws.max_column)
                cell.font = Font(bold=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for column in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[get_column_letter(column[0].column)].width = max_length + 2

        for col_idx, header in enumerate(headers, start=1):
            if header != "Data" and header != "Descrição":
                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row=row_idx, column=col_idx).number_format = '#,##0.00'

        wb.save(file_path)

        return str(file_path.resolve())
